"""Build the advisor briefing from its reviewed Markdown source; no model imports."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/advisor_briefing.md'
OUTPUT = ROOT / 'docs/advisor_briefing.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    s.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.bold = name.startswith('Heading')
    s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.15
    for border in s.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for style in doc.styles:
    for border in style.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Caption'].font.size = Pt(9.5)
doc.styles['Caption'].font.italic = False
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].paragraph_format.space_before = Pt(9)
    doc.styles[name].paragraph_format.space_after = Pt(7)
header = sec.header.paragraphs[0]
header.text = '扩散模型语义水印黑盒伪造研究　｜　导师讨论材料'
header.style = 'Caption'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026 年 10 月 16 日　 ·　')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for r in footer.runs: r.font.size = Pt(9)

def table(rows):
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = {3: [1.35, 2.45, 3.2], 4: [1.45, 2.0, 1.75, 1.8], 5: [1.5, 1.375, 1.375, 1.375, 1.375]}[len(rows[0])]
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i,row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j,(cell, txt) in enumerate(zip(cells,row)):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3); p.paragraph_format.line_spacing=1.08
            p.add_run(txt)
            if re.fullmatch(r'[\d.%± /]+',txt): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs: run.font.size=Pt(9.5); run.bold=(i==0)
            tcpr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF4');tcpr.append(shade)
        if i==0:
            rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=SOURCE.read_text(encoding='utf-8').splitlines();i=0;references=False
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line: continue
    if line=='<!-- PAGE -->': doc.add_page_break(); continue
    if line.startswith('|'):
        rows=[]
        while True:
            if not re.fullmatch(r'[| :\-]+',line): rows.append([v.strip() for v in line.strip('|').split('|')])
            if i>=len(lines) or not lines[i].startswith('|'):break
            line=lines[i];i+=1
        table(rows);continue
    if line.startswith('![PAIR]('):
        paths=line[8:-1].split('|');p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for j,path in enumerate(paths):
            p.add_run().add_picture(str((SOURCE.parent/path).resolve()),width=Inches(2.8))
            if j==0:p.add_run('  ')
        p.paragraph_format.keep_with_next=True
        continue
    if line.startswith('# '): doc.add_paragraph(line[2:],'Title');continue
    if line.startswith('## '):
        references=line.startswith('## 10 ')
        doc.add_paragraph(line[3:],'Heading 1');continue
    if line.startswith('### '): doc.add_paragraph(line[4:],'Heading 2');continue
    p=doc.add_paragraph(line, 'Caption' if line.startswith(('图 1','来源：','https://')) else None)
    p.paragraph_format.widow_control=True
    if references:
        p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.line_spacing=1.05
        for r in p.runs:r.font.size=Pt(10)
    if line.startswith('https://'):
        for r in p.runs:r.font.size=Pt(8.5)
doc.core_properties.title='扩散模型语义水印黑盒伪造研究讨论'
doc.core_properties.subject='基础论文 防御基线 研究问题 与评估契约'
doc.core_properties.author='Watermark Research'
doc.core_properties.keywords='semantic watermark; forgery; Reprompting; SemBind; geometry'
doc.save(OUTPUT)
print(OUTPUT)
